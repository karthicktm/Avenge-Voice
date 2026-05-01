"""Document processing service for text extraction and chunking."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import TYPE_CHECKING, ClassVar

import structlog
from charset_normalizer import from_bytes

from app.core.config import settings

if TYPE_CHECKING:
    from app.services.translation_service import TranslationService

logger = structlog.get_logger()


class DocumentProcessor:
    """Process documents for RAG: extract text and split into chunks.

    Supported file types:
    - PDF: Uses pypdf for extraction
    - DOCX: Uses python-docx for extraction
    - TXT: Direct text reading
    - MD: Direct text reading (Markdown)
    - XLSX: Uses openpyxl for extraction
    - XLS: Uses xlrd for extraction
    - CSV: Uses stdlib csv module
    """

    SUPPORTED_TYPES: ClassVar[set[str]] = {"pdf", "docx", "txt", "md", "xlsx", "xls", "csv"}

    def __init__(self) -> None:
        """Initialize document processor."""
        self.chunk_size = settings.RAG_CHUNK_SIZE
        self.chunk_overlap = settings.RAG_CHUNK_OVERLAP
        self.logger = logger.bind(
            component="document_processor",
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
        )

    def get_file_type(self, filename: str) -> str | None:
        """Get file type from filename.

        Args:
            filename: Filename with extension

        Returns:
            File type or None if unsupported
        """
        ext = Path(filename).suffix.lower().lstrip(".")
        return ext if ext in self.SUPPORTED_TYPES else None

    def _sanitize_text(self, text: str) -> str:
        """Sanitize text by removing null bytes and other problematic characters.

        PostgreSQL UTF-8 encoding doesn't accept null bytes (0x00).

        Args:
            text: Raw extracted text

        Returns:
            Sanitized text safe for database storage
        """
        # Remove null bytes that PostgreSQL can't handle
        text = text.replace("\x00", "")
        # Remove other control characters except whitespace (newline, tab, carriage return, space)
        allowed_whitespace = {"\n", "\t", "\r", " "}
        text = "".join(char for char in text if char in allowed_whitespace or char.isprintable())
        return text

    def _detect_encoding(self, content: bytes) -> str:
        """Detect the encoding of text content.

        Uses charset-normalizer for robust encoding detection,
        supporting international characters (Swedish å, ä, ö, etc.).

        Args:
            content: Raw file bytes

        Returns:
            Detected encoding name (defaults to utf-8)
        """
        # First try UTF-8 with BOM detection
        if content.startswith(b"\xef\xbb\xbf"):
            self.logger.debug("encoding_detected", encoding="utf-8-sig", method="bom")
            return "utf-8-sig"

        # Use charset-normalizer for detection
        result = from_bytes(content)
        best_match = result.best()

        if best_match is not None:
            encoding = best_match.encoding
            self.logger.info(
                "encoding_detected",
                encoding=encoding,
            )
            return encoding

        # Default to UTF-8 if detection fails
        self.logger.warning("encoding_detection_failed", fallback="utf-8")
        return "utf-8"

    async def extract_text(self, content: bytes, file_type: str) -> str:
        """Extract text from document content.

        Args:
            content: Raw file bytes
            file_type: File type (pdf, docx, txt, md)

        Returns:
            Extracted text content
        """
        if file_type == "pdf":
            text = await self._extract_pdf(content)
        elif file_type == "docx":
            text = await self._extract_docx(content)
        elif file_type in ("txt", "md"):
            text = await self._extract_text_file(content)
        elif file_type == "xlsx":
            text = await self._extract_xlsx(content)
        elif file_type == "xls":
            text = await self._extract_xls(content)
        elif file_type == "csv":
            text = await self._extract_csv(content)
        else:
            msg = f"Unsupported file type: {file_type}"
            raise ValueError(msg)

        # Sanitize text to remove null bytes and other problematic characters
        return self._sanitize_text(text)

    async def _extract_text_file(self, content: bytes) -> str:
        """Extract text from TXT/MD files with encoding detection.

        Supports multiple encodings including UTF-8, ISO-8859-1,
        Windows-1252, and other common encodings for international text.

        Args:
            content: Raw file bytes

        Returns:
            Extracted text
        """
        encoding = self._detect_encoding(content)

        try:
            text = content.decode(encoding)
            self.logger.info(
                "text_file_decoded",
                encoding=encoding,
                text_length=len(text),
            )
            return text
        except (UnicodeDecodeError, LookupError) as e:
            # Fallback: try UTF-8 with error replacement
            self.logger.warning(
                "encoding_fallback",
                original_encoding=encoding,
                error=str(e),
            )
            return content.decode("utf-8", errors="replace")

    async def _extract_pdf(self, content: bytes) -> str:
        """Extract text from PDF.

        Args:
            content: PDF file bytes

        Returns:
            Extracted text
        """
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            return "\n\n".join(text_parts)
        except Exception as e:
            self.logger.exception("pdf_extraction_failed")
            msg = f"Failed to extract PDF text: {e}"
            raise ValueError(msg) from e

    async def _extract_docx(self, content: bytes) -> str:
        """Extract text from DOCX.

        Args:
            content: DOCX file bytes

        Returns:
            Extracted text
        """
        try:
            from docx import Document

            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)
        except Exception as e:
            self.logger.exception("docx_extraction_failed")
            msg = f"Failed to extract DOCX text: {e}"
            raise ValueError(msg) from e

    async def _extract_xlsx(self, content: bytes) -> str:
        """Extract text from XLSX files.

        Args:
            content: XLSX file bytes

        Returns:
            Extracted text with sheet/row structure
        """
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sections: list[str] = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue
                headers = [str(h) if h is not None else "" for h in rows[0]]
                lines: list[str] = [f"Sheet: {sheet_name}"]
                for row in rows[1:]:
                    cells = [str(v) if v is not None else "" for v in row]
                    if not any(cells):
                        continue
                    parts = [
                        f"{headers[i]}: {cells[i]}"
                        for i in range(min(len(headers), len(cells)))
                        if headers[i]
                    ]
                    if parts:
                        lines.append(", ".join(parts))
                sections.append("\n".join(lines))
            wb.close()
            return "\n\n".join(sections)
        except Exception as e:
            self.logger.exception("xlsx_extraction_failed")
            msg = f"Failed to extract XLSX text: {e}"
            raise ValueError(msg) from e

    async def _extract_xls(self, content: bytes) -> str:
        """Extract text from XLS (legacy Excel) files.

        Args:
            content: XLS file bytes

        Returns:
            Extracted text with sheet/row structure
        """
        try:
            import xlrd

            wb = xlrd.open_workbook(file_contents=content)
            sections: list[str] = []
            for sheet_name in wb.sheet_names():
                ws = wb.sheet_by_name(sheet_name)
                if ws.nrows == 0:
                    continue
                headers = [str(ws.cell_value(0, col)) for col in range(ws.ncols)]
                lines: list[str] = [f"Sheet: {sheet_name}"]
                for row_idx in range(1, ws.nrows):
                    cells = [str(ws.cell_value(row_idx, col)) for col in range(ws.ncols)]
                    if not any(cells):
                        continue
                    parts = [f"{headers[i]}: {cells[i]}" for i in range(len(headers)) if headers[i]]
                    if parts:
                        lines.append(", ".join(parts))
                sections.append("\n".join(lines))
            return "\n\n".join(sections)
        except Exception as e:
            self.logger.exception("xls_extraction_failed")
            msg = f"Failed to extract XLS text: {e}"
            raise ValueError(msg) from e

    async def _extract_csv(self, content: bytes) -> str:
        """Extract text from CSV files.

        Args:
            content: CSV file bytes

        Returns:
            Extracted text with header-prefixed rows
        """
        try:
            encoding = self._detect_encoding(content)
            try:
                text = content.decode(encoding)
            except (UnicodeDecodeError, LookupError):
                text = content.decode("utf-8", errors="replace")

            reader = csv.reader(io.StringIO(text))
            rows = list(reader)
            if not rows:
                return ""
            headers = rows[0]
            lines: list[str] = []
            for row in rows[1:]:
                if not any(row):
                    continue
                parts = [
                    f"{headers[i]}: {row[i]}"
                    for i in range(min(len(headers), len(row)))
                    if headers[i]
                ]
                if parts:
                    lines.append(", ".join(parts))
            return "\n".join(lines)
        except Exception as e:
            self.logger.exception("csv_extraction_failed")
            msg = f"Failed to extract CSV text: {e}"
            raise ValueError(msg) from e

    def chunk_text(self, text: str) -> list[str]:
        """Split text into overlapping chunks.

        Uses character-based chunking with overlap to ensure
        context continuity across chunks.

        Args:
            text: Full text content

        Returns:
            List of text chunks
        """
        if not text or not text.strip():
            return []

        # Clean and normalize whitespace
        text = " ".join(text.split())

        chunks = []
        start = 0

        while start < len(text):
            # Calculate end position
            end = start + self.chunk_size

            if end >= len(text):
                # Last chunk
                chunk = text[start:].strip()
                if chunk:
                    chunks.append(chunk)
                break

            # Try to break at sentence boundary
            chunk_text = text[start:end]
            last_period = chunk_text.rfind(". ")
            last_newline = chunk_text.rfind("\n")
            last_question = chunk_text.rfind("? ")
            last_exclaim = chunk_text.rfind("! ")

            # Find best break point
            break_point = max(last_period, last_newline, last_question, last_exclaim)

            if break_point > self.chunk_size // 2:
                # Found a good break point in second half of chunk
                end = start + break_point + 1

            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)

            # Move start with overlap
            start = end - self.chunk_overlap

        self.logger.info(
            "text_chunked",
            original_length=len(text),
            chunk_count=len(chunks),
            avg_chunk_size=len(text) // len(chunks) if chunks else 0,
        )

        return chunks

    async def process(self, content: bytes, filename: str) -> tuple[str, list[str]]:
        """Process document: extract text and create chunks.

        Args:
            content: Raw file bytes
            filename: Original filename

        Returns:
            Tuple of (full_text, list_of_chunks)
        """
        file_type = self.get_file_type(filename)
        if not file_type:
            msg = f"Unsupported file type for {filename}"
            raise ValueError(msg)

        self.logger.info(
            "processing_document",
            filename=filename,
            file_type=file_type,
            content_size=len(content),
        )

        # Extract text
        full_text = await self.extract_text(content, file_type)

        if not full_text or not full_text.strip():
            msg = f"No text content extracted from {filename}"
            raise ValueError(msg)

        # Create chunks
        chunks = self.chunk_text(full_text)

        self.logger.info(
            "document_processed",
            filename=filename,
            text_length=len(full_text),
            chunk_count=len(chunks),
        )

        return full_text, chunks

    async def process_with_translation(
        self,
        content: bytes,
        filename: str,
        translation_service: TranslationService | None = None,
    ) -> tuple[str, list[str], list[str], str, str]:
        """Process document with optional translation for cross-lingual support.

        Extracts text, chunks it, detects language, and optionally translates
        to English for embedding. Original content is preserved for display.

        Args:
            content: Raw file bytes
            filename: Original filename
            translation_service: Optional translation service for cross-lingual support

        Returns:
            Tuple of:
                - full_text: Complete extracted text
                - original_chunks: List of chunks in original language
                - translated_chunks: List of chunks translated to English (or original if no translation)
                - source_language: Detected language code (e.g., "sv", "en")
                - translation_status: "not_needed", "completed", "partial", "failed", or "skipped"
        """
        # First, process normally to get text and chunks
        full_text, original_chunks = await self.process(content, filename)

        # If no translation service, skip translation
        if translation_service is None:
            return full_text, original_chunks, original_chunks, "en", "skipped"

        # Detect language
        source_language = await translation_service.detect_language(full_text)

        self.logger.info(
            "language_detected",
            filename=filename,
            source_language=source_language,
        )

        # If already English, no translation needed
        if source_language == "en":
            return full_text, original_chunks, original_chunks, "en", "not_needed"

        # Translate chunks
        translated_chunks, status = await translation_service.translate_chunks(
            original_chunks, source_language
        )

        self.logger.info(
            "document_translated",
            filename=filename,
            source_language=source_language,
            status=status,
            original_chunk_count=len(original_chunks),
            translated_chunk_count=len(translated_chunks),
        )

        return full_text, original_chunks, translated_chunks, source_language, status
